import csv
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
import re
from time import perf_counter
from typing import Any
from zipfile import BadZipFile

from fastapi import UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import delete, desc, func, select
from sqlalchemy.orm import Session, selectinload

from app.core.config import settings
from app.core.exceptions import AnalysisFileError, ResourceNotFoundError
from app.database.models import (
    AnalysisRun,
    AnalysisValidationError,
    Finding,
    RuleExecution,
    Transaction,
    utc_now,
)
from app.rules.context import RuleContext, normalize_key
from app.rules.engine import rule_engine
from app.services.rules import rule_service
from app.schemas.analysis import (
    AnalysisDetail,
    AnalysisHistoryItem,
    AnalysisResponse,
    DashboardSummary,
    FindingResponse,
    FindingSummary,
    RowError,
    RuleExecutionResponse,
    TransactionResponse,
)
from app.validators.analysis import validate_columns, validate_row

ALLOWED_EXTENSIONS = {".csv", ".xlsx"}


class AnalysisService:
    async def process_upload(self, file: UploadFile | None, db: Session) -> AnalysisResponse:
        if file is None or not file.filename:
            raise AnalysisFileError("file_required", "Please select a file.")
        filename = Path(file.filename).name
        extension = Path(filename).suffix.lower()
        if extension not in ALLOWED_EXTENSIONS:
            raise AnalysisFileError(
                "invalid_extension", "Unsupported format. Please use a CSV or XLSX file."
            )

        started = perf_counter()
        content = await file.read(settings.noams_max_upload_bytes + 1)
        if len(content) > settings.noams_max_upload_bytes:
            raise AnalysisFileError(
                "file_too_large",
                f"The file exceeds the {settings.noams_max_upload_bytes // 1_048_576} MB limit.",
                413,
            )
        if not content:
            raise AnalysisFileError("empty_file", "The file is empty.")

        columns, rows = self._read_file(content, extension)
        self._ensure_columns(columns)
        if not rows:
            raise AnalysisFileError("empty_file", "The file does not contain any records.")

        row_errors: list[RowError] = []
        valid_rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, row in enumerate(rows, start=2):
            errors = validate_row(row, row_number)
            if errors:
                row_errors.extend(errors)
            else:
                valid_rows.append((row_number, row))

        run = AnalysisRun(
            filename=filename,
            source_type=extension.removeprefix("."),
            status="validating",
            total_records=len(rows),
            valid_records=len(valid_rows),
            invalid_records=len(rows) - len(valid_rows),
            started_at=utc_now(),
        )
        stored_file: Path | None = None
        try:
            db.add(run)
            db.flush()
            db.add_all(
                AnalysisValidationError(
                    analysis_run_id=run.id,
                    row=error.row,
                    column=error.column,
                    message=error.message,
                )
                for error in row_errors
            )
            transactions = [self._normalize_row(run.id, number, row) for number, row in valid_rows]
            db.add_all(transactions)
            db.flush()
            run.status = "processing"
            rule_configurations = rule_service.execution_configurations(db)

            context = RuleContext(
                analysis_run=run,
                transactions=transactions,
                historical_amounts=self._historical_amounts(db, run, transactions),
                settings=settings,
                session=db,
                rule_configurations=rule_configurations,
            )
            findings, executions = rule_engine.execute(context)
            run.status = "completed_with_findings" if findings else "completed"
            run.completed_at = utc_now()
            run.processing_time_ms = max(1, round((perf_counter() - started) * 1000))
            stored_file = self._store_upload(content, run.run_uuid, filename, extension)
            run.stored_file_path = str(stored_file)
            run.file_size_bytes = len(content)
            db.commit()

            transaction_map = {transaction.id: transaction for transaction in transactions}
            findings.sort(key=self._finding_sort_key)
            return AnalysisResponse(
                analysis_id=run.id,
                filename=run.filename,
                status=run.status,
                total_records=run.total_records,
                valid_records=run.valid_records,
                invalid_records=run.invalid_records,
                processing_time_ms=run.processing_time_ms,
                errors=row_errors,
                findings_summary=self._summarize(findings),
                findings=[self.finding_schema(item, transaction_map.get(item.transaction_id)) for item in findings],
                rule_executions=[RuleExecutionResponse.model_validate(item) for item in executions],
            )
        except AnalysisFileError:
            db.rollback()
            self._remove_stored_file(stored_file)
            raise
        except Exception as exc:
            db.rollback()
            self._remove_stored_file(stored_file)
            self._record_failed_run(db, run, started, exc)
            raise AnalysisFileError(
                "analysis_failed", "The file analysis could not be completed.", 500
            ) from exc

    def dashboard_summary(self, db: Session) -> DashboardSummary:
        analyses_count, records_processed = db.execute(
            select(func.count(AnalysisRun.id), func.coalesce(func.sum(AnalysisRun.total_records), 0))
        ).one()
        findings_count = db.scalar(select(func.count(Finding.id))) or 0
        critical_risks = db.scalar(
            select(func.count(Finding.id)).where(Finding.severity == "critical", Finding.status == "open")
        ) or 0
        open_findings = db.scalars(select(Finding).where(Finding.status == "open")).all()
        # One maximum exposure per transaction avoids stacking multiple rules on the same purchase.
        linked: dict[int, Decimal] = {}
        aggregate = Decimal("0")
        for finding in open_findings:
            impact = finding.estimated_impact or Decimal("0")
            if finding.transaction_id is None:
                aggregate += impact
            else:
                linked[finding.transaction_id] = max(linked.get(finding.transaction_id, Decimal("0")), impact)
        last_run = db.scalar(select(AnalysisRun).order_by(desc(AnalysisRun.created_at)).limit(1))
        last_run_at = last_run.created_at if last_run else None
        if last_run_at and last_run_at.tzinfo is None:
            last_run_at = last_run_at.replace(tzinfo=timezone.utc)
        return DashboardSummary(
            analyses_count=analyses_count,
            records_processed=records_processed,
            findings_count=findings_count,
            critical_risks=critical_risks,
            estimated_impact=sum(linked.values(), aggregate),
            last_run_at=last_run_at,
            last_run_status=last_run.status if last_run else None,
        )

    def list_analyses(
        self,
        db: Session,
        *,
        limit: int = 50,
        offset: int = 0,
        filename: str | None = None,
        status: str | None = None,
    ) -> list[AnalysisHistoryItem]:
        grouped_runs = self._grouped_analysis_runs()
        query = select(grouped_runs).where(grouped_runs.c.version_rank == 1)
        if filename:
            query = query.where(
                func.lower(grouped_runs.c.filename).contains(filename.casefold())
            )
        if status:
            query = query.where(grouped_runs.c.status == status)
        rows = db.execute(
            query.order_by(desc(grouped_runs.c.created_at), desc(grouped_runs.c.id))
            .offset(offset)
            .limit(limit)
        ).mappings().all()
        return [
            AnalysisHistoryItem(
                **row,
                is_updated=row["versions_count"] > 1,
            )
            for row in rows
        ]

    @staticmethod
    def _grouped_analysis_runs():
        """Select the latest stored execution and total versions for each exact filename."""
        return select(
            AnalysisRun.id,
            AnalysisRun.filename,
            AnalysisRun.created_at,
            AnalysisRun.status,
            AnalysisRun.total_records,
            AnalysisRun.valid_records,
            AnalysisRun.invalid_records,
            AnalysisRun.findings_count,
            AnalysisRun.processing_time_ms,
            AnalysisRun.source_type,
            func.count(AnalysisRun.id).over(
                partition_by=AnalysisRun.filename
            ).label("versions_count"),
            func.row_number().over(
                partition_by=AnalysisRun.filename,
                order_by=(desc(AnalysisRun.created_at), desc(AnalysisRun.id)),
            ).label("version_rank"),
        ).subquery()

    def get_analysis(self, db: Session, analysis_id: int) -> AnalysisDetail:
        run = db.scalar(
            select(AnalysisRun)
            .where(AnalysisRun.id == analysis_id)
            .options(
                selectinload(AnalysisRun.transactions),
                selectinload(AnalysisRun.findings).selectinload(Finding.transaction),
                selectinload(AnalysisRun.rule_executions),
                selectinload(AnalysisRun.validation_errors),
            )
        )
        if run is None:
            raise ResourceNotFoundError("analysis_not_found", "The requested analysis does not exist.")
        findings = sorted(run.findings, key=self._finding_sort_key)
        return AnalysisDetail(
            analysis_id=run.id,
            run_uuid=run.run_uuid,
            filename=run.filename,
            source_type=run.source_type,
            status=run.status,
            total_records=run.total_records,
            valid_records=run.valid_records,
            invalid_records=run.invalid_records,
            processing_time_ms=run.processing_time_ms,
            errors=[
                RowError(row=error.row, column=error.column, message=error.message)
                for error in sorted(run.validation_errors, key=lambda item: (item.row, item.id))
            ],
            findings_summary=self._summarize(findings),
            findings=[self.finding_schema(item, item.transaction) for item in findings],
            rule_executions=[RuleExecutionResponse.model_validate(item) for item in run.rule_executions],
            started_at=run.started_at,
            completed_at=run.completed_at,
            transactions=[TransactionResponse.model_validate(item) for item in run.transactions],
        )

    def delete_analysis(self, db: Session, analysis_id: int) -> None:
        run = db.get(AnalysisRun, analysis_id)
        if run is None:
            raise ResourceNotFoundError(
                "analysis_not_found", "The requested analysis does not exist."
            )

        stored_file = Path(run.stored_file_path) if run.stored_file_path else None
        try:
            db.execute(delete(Finding).where(Finding.analysis_run_id == analysis_id))
            db.execute(delete(RuleExecution).where(RuleExecution.analysis_run_id == analysis_id))
            db.execute(
                delete(AnalysisValidationError).where(
                    AnalysisValidationError.analysis_run_id == analysis_id
                )
            )
            db.execute(delete(Transaction).where(Transaction.analysis_run_id == analysis_id))
            db.delete(run)
            db.commit()
        except Exception:
            db.rollback()
            raise

        self._remove_stored_file(stored_file)

    def _normalize_row(self, run_id: int, row_number: int, row: dict[str, Any]) -> Transaction:
        return Transaction(
            analysis_run_id=run_id,
            transaction_id=self._clean(row["transaction_id"]),
            transaction_date=self._date(row["transaction_date"]),
            supplier_name=self._clean(row["supplier_name"]),
            department=self._clean(row["department"]),
            category=self._clean(row["category"]),
            quantity=Decimal(str(row["quantity"]).strip()),
            unit_price=Decimal(str(row["unit_price"]).strip()),
            total_amount=Decimal(str(row["total_amount"]).strip()),
            approval_status=self._clean(row["approval_status"]).casefold(),
            purchase_order=self._clean(row.get("purchase_order")) or None,
            budget_amount=(
                Decimal(str(row["budget_amount"]).strip())
                if row.get("budget_amount") not in (None, "")
                else None
            ),
            currency=(self._clean(row.get("currency")) or "PEN").upper()[:3],
            source_row_number=row_number,
        )

    def _historical_amounts(
        self, db: Session, run: AnalysisRun, transactions: list[Transaction]
    ) -> dict[str, list[Decimal]]:
        supplier_keys = {normalize_key(item.supplier_name) for item in transactions}
        history = db.execute(
            select(Transaction.supplier_name, Transaction.total_amount)
            .join(AnalysisRun)
            .where(
                AnalysisRun.id < run.id,
                AnalysisRun.status.in_(("completed", "completed_with_findings")),
            )
            .order_by(Transaction.id.desc())
            .limit(10_000)
        ).all()
        result: dict[str, list[Decimal]] = {key: [] for key in supplier_keys}
        for supplier_name, amount in history:
            key = normalize_key(supplier_name)
            if key in result:
                result[key].append(amount)
        return result

    def _summarize(self, findings: list[Finding]) -> FindingSummary:
        counts = Counter(item.severity for item in findings)
        return FindingSummary(
            total=len(findings),
            low=counts["low"],
            medium=counts["medium"],
            high=counts["high"],
            critical=counts["critical"],
            estimated_impact=sum(
                (item.estimated_impact or Decimal("0") for item in findings), Decimal("0")
            ),
        )

    def finding_schema(self, finding: Finding, transaction: Transaction | None) -> FindingResponse:
        return FindingResponse(
            id=finding.id,
            finding_uuid=finding.finding_uuid,
            rule_code=finding.rule_code,
            title=finding.title,
            description=finding.description,
            category=finding.category,
            severity=finding.severity,
            status=finding.status,
            detected_value=finding.detected_value,
            expected_value=finding.expected_value,
            deviation_percentage=finding.deviation_percentage,
            estimated_impact=finding.estimated_impact,
            confidence_score=finding.confidence_score,
            transaction_id=finding.transaction_id,
            detected_at=finding.detected_at,
            evidence_json=finding.evidence_json,
            supplier_name=transaction.supplier_name if transaction else None,
            transaction_date=transaction.transaction_date if transaction else None,
            purchase_order=transaction.purchase_order if transaction else None,
        )

    @staticmethod
    def _finding_sort_key(finding: Finding):
        order = {"critical": 0, "high": 1, "medium": 2, "low": 3}
        return (order.get(finding.severity, 4), -(finding.estimated_impact or Decimal("0")))

    def _record_failed_run(
        self, db: Session, attempted: AnalysisRun, started: float, exc: Exception
    ) -> None:
        try:
            failed = AnalysisRun(
                filename=attempted.filename,
                source_type=attempted.source_type,
                status="failed",
                total_records=attempted.total_records,
                valid_records=attempted.valid_records,
                invalid_records=attempted.invalid_records,
                processing_time_ms=max(1, round((perf_counter() - started) * 1000)),
                error_message=f"{type(exc).__name__}: {str(exc)[:300]}",
                started_at=attempted.started_at,
                completed_at=utc_now(),
            )
            db.add(failed)
            db.commit()
        except Exception:
            db.rollback()

    @staticmethod
    def _store_upload(content: bytes, run_uuid: str, filename: str, extension: str) -> Path:
        upload_root = settings.noams_upload_dir.resolve()
        safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).stem).strip("._")
        safe_stem = (safe_stem or "upload")[:180]
        target = (upload_root / f"{run_uuid}_{safe_stem}{extension}").resolve()
        if target.parent != upload_root:
            raise AnalysisFileError("file_storage_error", "The uploaded file could not be stored.", 500)
        try:
            upload_root.mkdir(parents=True, exist_ok=True)
            target.write_bytes(content)
        except OSError as exc:
            try:
                target.unlink(missing_ok=True)
            except OSError:
                pass
            raise AnalysisFileError(
                "file_storage_error", "The uploaded file could not be stored.", 500
            ) from exc
        return target

    @staticmethod
    def _remove_stored_file(path: Path | None) -> None:
        if path is None:
            return
        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass

    @staticmethod
    def _clean(value: Any) -> str:
        return " ".join(str(value).strip().split()) if value is not None else ""

    @staticmethod
    def _date(value: Any) -> date:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        candidate = str(value).strip()
        try:
            return datetime.fromisoformat(candidate.replace("Z", "+00:00")).date()
        except ValueError:
            for date_format in ("%d/%m/%Y", "%m/%d/%Y"):
                try:
                    return datetime.strptime(candidate, date_format).date()
                except ValueError:
                    continue
        raise ValueError("Invalid validated date")

    def _read_file(self, content: bytes, extension: str) -> tuple[list[str], list[dict[str, Any]]]:
        return self._read_csv(content) if extension == ".csv" else self._read_xlsx(content)

    def _read_csv(self, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
        try:
            reader = csv.DictReader(StringIO(content.decode("utf-8-sig")))
            if reader.fieldnames is None:
                raise AnalysisFileError("empty_file", "The CSV file is empty.")
            columns = [str(column).strip() for column in reader.fieldnames]
            rows = [
                {str(key).strip(): value for key, value in row.items() if key is not None}
                for row in reader
            ]
            return columns, rows
        except UnicodeDecodeError as exc:
            raise AnalysisFileError(
                "csv_read_error", "The CSV file could not be read. Please use UTF-8 encoding."
            ) from exc
        except csv.Error as exc:
            raise AnalysisFileError("csv_read_error", "The CSV file could not be read.") from exc

    def _read_xlsx(self, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, None)
            if header is None:
                raise AnalysisFileError("empty_file", "The Excel file is empty.")
            columns = [str(value).strip() if value is not None else "" for value in header]
            rows = [dict(zip(columns, values, strict=False)) for values in iterator]
            workbook.close()
            return columns, rows
        except AnalysisFileError:
            raise
        except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError) as exc:
            raise AnalysisFileError("xlsx_read_error", "The Excel file could not be read.") from exc

    def _ensure_columns(self, columns: list[str]) -> None:
        missing, unexpected = validate_columns(columns)
        if not missing and not unexpected and len(columns) == len(set(columns)):
            return
        details = []
        if missing:
            details.append(f"missing: {', '.join(missing)}")
        if unexpected:
            details.append(f"unsupported: {', '.join(unexpected)}")
        if len(columns) != len(set(columns)):
            details.append("duplicate columns detected")
        raise AnalysisFileError(
            "invalid_columns", f"The columns do not match the required format ({'; '.join(details)})."
        )


analysis_service = AnalysisService()
